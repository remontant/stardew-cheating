"""에글이 님 스타듀밸리 공략 시트(.xlsx) -> data/data.js

시트가 업데이트되면 구글 시트에서 [파일 > 다운로드 > Microsoft Excel(.xlsx)] 로 받은 뒤
    pip install openpyxl
    python tools/extract.py 받은파일.xlsx
를 돌리면 data/data.js 가 새로 만들어진다.

시트는 사람이 보기 좋게 칸을 병합하고 색으로 의미를 나눈 "그림 같은 표"라서,
열 위치·배경색을 직접 짚어서 읽는다. 원작자가 표 모양을 크게 바꾸면 이 파일도 손봐야 한다.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import column_index_from_string as ci

SEASONS = ["봄", "여름", "가을", "겨울"]

# 시트 안의 오타. 화면에서 이름으로 서로 연결(생일 -> 선물)하므로 맞춰 둔다.
NAME_FIX = {"에블린린": "에블린", "하빈": "하비"}


def text(v):
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        v = "".join(str(p.text if hasattr(p, "text") else p) for p in v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("​", "").replace("\r", "")
    s = "\n".join(line.strip() for line in s.split("\n"))
    return re.sub(r"\n{2,}", "\n", s).strip()


def num(v):
    s = text(v).replace(",", "")
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return s or None


def fill_key(cell):
    """배경색. 테마색은 rgb 대신 theme/tint 로 구분한다(번들 범례가 테마색을 쓴다)."""
    f = cell.fill
    if not f or f.fill_type != "solid":
        return None
    c = f.fgColor
    if c.type == "rgb" and isinstance(c.rgb, str):
        return None if c.rgb in ("00000000", "FFFFFFFF") else c.rgb
    if c.type == "theme":
        return f"theme{c.theme}:{round(c.tint or 0, 3)}"
    return None


class Sheet:
    """병합 칸을 펼쳐 읽을 수 있게 감싼다. raw() 는 병합의 첫 칸에만 값이 있고, val() 은 병합 전체에 같은 값."""

    def __init__(self, ws):
        self.ws = ws
        self.anchor = {}
        self.span = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.anchor[(r, c)] = (rng.min_row, rng.min_col)
            self.span[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)

    def _c(self, c):
        return ci(c) if isinstance(c, str) else c

    def cell(self, r, c):
        return self.ws.cell(r, self._c(c))

    def raw(self, r, c):
        return text(self.cell(r, c).value)

    def val(self, r, c):
        r0, c0 = self.anchor.get((r, self._c(c)), (r, self._c(c)))
        return text(self.ws.cell(r0, c0).value)

    def fill(self, r, c):
        return fill_key(self.cell(r, c))

    def link(self, r, c):
        h = self.cell(r, c).hyperlink
        return h.target if h else None

    def spans_cols(self, r, c, c_end):
        r0, c0 = self.anchor.get((r, self._c(c)), (r, self._c(c)))
        _, c1 = self.span.get((r0, c0), (r0, c0))
        return c0 <= self._c(c) and c1 >= self._c(c_end)


def lines(s):
    return [x for x in (s or "").split("\n") if x]


# ---------------------------------------------------------------- 기본
def basics(sh):
    tips = []
    for r in range(2, 13):
        t = sh.raw(r, "A")
        if t:
            tips.append(t)
    return {"tips": tips, "video": sh.link(1, "A")}


# ---------------------------------------------------------------- 선물
def gift_items(v):
    """빨간 글씨 = 구하기 쉬운 추천 선물. 쉼표·줄바꿈으로 나누고, 괄호만 있는 조각은 앞 항목의 메모로 붙인다."""
    parts = []
    if isinstance(v, CellRichText):
        for p in v:
            t = str(p.text if hasattr(p, "text") else p)
            red = False
            if hasattr(p, "font") and p.font is not None and p.font.color is not None:
                red = isinstance(p.font.color.rgb, str) and p.font.color.rgb.upper() == "FFFF0000"
            parts.append((t, red))
    else:
        parts.append((text(v), False))
    items = []
    for t, red in parts:
        for piece in re.split(r"[,\n]", t):
            piece = piece.replace("​", "").strip()
            if not piece:
                continue
            if piece.startswith("(") and items:
                items[-1]["memo"] = (items[-1].get("memo", "") + " " + piece).strip()
                continue
            m = re.match(r"^(.*?)\s*(\(.*\))$", piece)
            item = {"name": piece, "rec": red}
            if m and m.group(1):
                item = {"name": m.group(1).strip(), "rec": red, "memo": m.group(2)}
            items.append(item)
    return items


def parse_birthday(s):
    m = re.match(r"(봄|여름|가을|겨울)\s*(\d+)", s)
    return {"season": m.group(1), "day": int(m.group(2))} if m else None


def villagers(sh):
    out = []
    for r in range(8, 40):
        for cn, cb, cg, marry in (("B", "C", "D", True), ("F", "G", "H", False)):
            name = sh.raw(r, cn)
            if not name:
                continue
            name = NAME_FIX.get(name, name)
            out.append({
                "name": name,
                "marry": marry,
                "birthday": parse_birthday(sh.raw(r, cb)),
                "gifts": gift_items(sh.cell(r, cg).value),
            })
    notes = [sh.raw(r, "A") for r in range(2, 5) if sh.raw(r, "A")]
    return {"list": out, "notes": notes, "video": sh.link(1, "A")}


# ---------------------------------------------------------------- 달력
def calendar(sh, vlist):
    """생일은 선물 시트의 생일 칸을 기준으로 다시 채운다. 달력 시트는 가을 5일에 엘리엇 대신 알렉스가 적혀 있다."""
    names = {v["name"] for v in vlist}
    days = {s: {} for s in SEASONS}
    for r in range(1, sh.ws.max_row + 1):
        season = sh.raw(r, "C")
        if season not in SEASONS:
            continue
        for week in range(4):
            rr = r + 2 + week
            for i in range(7):
                dc = 2 + i * 2
                d = num(sh.raw(rr, dc))
                if not isinstance(d, int):
                    continue
                births, events = [], []
                for ln in lines(sh.raw(rr, dc + 1)):
                    ln = NAME_FIX.get(ln, ln)
                    (births if ln in names else events).append(ln)
                days[season][d] = {"birthdays": births, "events": events}
    for s in SEASONS:
        for d in days[s].values():
            d["birthdays"] = []
    for v in vlist:
        b = v["birthday"]
        if b:
            days[b["season"]].setdefault(b["day"], {"birthdays": [], "events": []})["birthdays"].append(v["name"])
    notes = [sh.raw(r, "A") for r in range(2, 4) if sh.raw(r, "A")]
    return {"days": days, "notes": notes, "video": sh.link(1, "A")}


# ---------------------------------------------------------------- 마을회관 번들
# 범례 칸의 꾸러미 이름 -> 마을회관 방. 시트에는 방 구분이 없어 게임 기준으로 붙였다.
ROOMS = [
    ("식료품 저장실", ["봄 작물", "여름 작물", "가을 작물", "고품질 작물", "동물", "장인"]),
    ("공예실", ["봄 채집", "여름 채집", "가을 채집", "겨울 채집", "건축", "이국적인 채집"]),
    ("어항", ["강 물고기", "호수 물고기", "바다 물고기", "밤 물고기", "게잡이 통발", "특별한 물고기"]),
    ("보일러실", ["대장장이", "지질학자", "모험가"]),
    ("게시판", ["요리사", "염료", "현장 연구", "사료", "마법사"]),
]


LABEL_FIX = {"야채고추냉이": "야생고추냉이"}


def bundles(sh):
    legend = {}
    order = []
    for r in range(10, 41):
        t = sh.raw(r, "A")
        k = sh.fill(r, "A")
        if not t or not k:
            continue
        m = re.match(r"^(.*?꾸러미)\s*(\d+)개 선택$", t)
        name, pick = (m.group(1), int(m.group(2))) if m else (t, None)
        legend[k] = name
        order.append({"name": name, "pick": pick, "items": []})
    by_name = {b["name"]: b for b in order}

    unknown = []

    def add(r, ic, cat, where, season):
        label = sh.raw(r, ic)
        label = LABEL_FIX.get(label, label)
        if not label:
            return
        k = sh.fill(r, ic)
        b = legend.get(k)
        if not b:
            unknown.append((r, ic, label, k))
            return
        items = by_name[b]["items"]
        for it in items:
            if it["label"] == label:
                if not season and not it["seasons"]:
                    it["times"] = it.get("times", 1) + 1  # 건축 꾸러미의 "나무 99개" 처럼 같은 칸이 두 번
                if season and season not in it["seasons"]:
                    it["seasons"].append(season)
                if where and where not in it["where"]:
                    it["where"].append(where)
                return
        items.append({"label": label, "cat": cat, "seasons": [season] if season else [], "where": [where] if where else []})

    # 계절 블록: 꾸러미 칸은 D/H/L/P, 그 왼쪽이 분류(채집·작물·낚시…), 오른쪽이 위치 메모
    for season, cc in zip(SEASONS, ("C", "G", "K", "O")):
        c = ci(cc)
        for r in range(10, 41):
            add(r, c + 1, sh.val(r, c), sh.val(r, c + 2), season)
    # 아래쪽 블록(계절 무관): 동물/가공 · 그 외 채집 · 그 외 낚시 · 광산
    for cc in ("C", "G", "K", "O"):
        c = ci(cc)
        group = sh.raw(45, c)
        for r in range(46, 64):
            sub = sh.val(r, c)
            add(r, c + 1, f"{group} · {sub}" if sub else group, sh.val(r, c + 2), None)

    if unknown:
        print("  ! 꾸러미 색을 못 찾은 칸:", unknown, file=sys.stderr)

    rooms = []
    for room, keys in ROOMS:
        bs = []
        for key in keys:
            b = next((x for x in order if x["name"].startswith(key + " ") or x["name"] == key + " 꾸러미"), None)
            if b:
                bs.append(b)
        rooms.append({"room": room, "bundles": bs})
    rooms.append({"room": "금고", "note": sh.raw(41, "A"), "bundles": [
        {"name": f"{g} 골드", "pick": None, "items": [{"label": f"{g}골드", "cat": "골드", "seasons": [], "where": []}]}
        for g in ("2,500", "5,000", "10,000", "25,000")
    ]})
    placed = {b["name"] for rm in rooms for b in rm["bundles"]}
    left = [b for b in order if b["name"] not in placed]
    if left:
        rooms.append({"room": "기타", "bundles": left})
    notes = [sh.raw(r, "A") for r in range(2, 7) if sh.raw(r, "A")]
    return {"rooms": rooms, "notes": notes, "video": sh.link(1, "A")}


# ---------------------------------------------------------------- 작물
GROW_COLS = ["F", "G", "H", "I", "J", "K", "L", "M"]
SELL_COLS = ["N", "O", "P", "Q", "R", "S", "T", "U"]
CAT_FIX = {"얜뭐지": "기타"}


def crops_sheet(sh, season, out):
    # 머리줄은 "비고" 칸으로 찾는다. 씨앗 획득처 값 중에 "씨앗 위치" 가 있어 "씨앗" 으로 찾으면 틀린다.
    header_rows = [r for r in range(1, 60) if sh.raw(r, "D") == "비고"]
    for hi, h in enumerate(header_rows):
        end = header_rows[hi + 1] if hi + 1 < len(header_rows) else h + 30
        top = sh.val(h, "V")
        if "벌꿀" in top:
            proc_labels = ["벌꿀", "벌꿀술(술통)", None]
        else:
            proc_labels = [sh.val(h + 1, c) or None for c in ("V", "X", "Z")]
        for r in range(h + 3, end):
            name = sh.raw(r, "B")
            if not name:
                continue
            name_lines = lines(name)
            cname = name_lines[0]
            memo = " ".join(name_lines[1:])
            procs = []
            for label, c in zip(proc_labels, ("V", "X", "Z")):
                base = num(sh.raw(r, c))
                art = num(sh.raw(r, ci(c) + 1))
                if base is None:
                    continue
                if isinstance(base, str) and not label:
                    procs.append({"label": None, "note": base})
                elif isinstance(base, str):
                    procs.append({"label": label, "note": base})
                else:
                    procs.append({"label": label, "base": base, "artisan": art if isinstance(art, (int, float)) else None})
            crop = {
                "name": cname,
                "memo": memo,
                "cat": CAT_FIX.get(sh.val(r, "A"), sh.val(r, "A")),
                "seasons": [season],
                "source": sh.raw(r, "C"),
                "notes": lines(sh.raw(r, "D")),
                "seedPrice": num(sh.raw(r, "E")),
                "grow": [num(sh.raw(r, c)) for c in GROW_COLS],
                "sell": [num(sh.raw(r, c)) for c in SELL_COLS],
                "proc": procs,
            }
            prev = next((x for x in out if x["name"] == cname), None)
            if prev:
                if season not in prev["seasons"]:
                    prev["seasons"].append(season)
            else:
                out.append(crop)


def crops(wb):
    out = []
    for title, season in (("봄 작물", "봄"), ("여름 작물", "여름"), ("가을 작물", "가을"), ("그 외 작물", "그 외")):
        crops_sheet(Sheet(wb[title]), season, out)
    sh = Sheet(wb["봄 작물"])
    notes = [sh.raw(r, "C") for r in range(2, 5) if sh.raw(r, "C")]
    return {"list": out, "notes": notes, "video": sh.link(1, "C")}


# ---------------------------------------------------------------- 채집
def forage(sh):
    out = []
    for r in range(28, 60):
        if sh.raw(r, "C"):
            g = sh.val(r, "B")
            out.append({"group": "기타" if g.startswith("?") else g, "name": sh.raw(r, "C"),
                        "where": sh.raw(r, "D"), "notes": lines(sh.raw(r, "E"))})
    for r in range(28, 60):
        if sh.raw(r, "H"):
            out.append({"group": sh.val(r, "G"), "name": sh.raw(r, "H"),
                        "where": sh.raw(r, "I"), "notes": lines(sh.raw(r, "J"))})
    return {"list": out}


# ---------------------------------------------------------------- 낚시
def fix_time(t):
    t = re.sub(r"(\d{2})\.(\d{2})", r"\1:\2", t)
    t = re.sub(r"(?<![\d:])(\d{2})(\d{2})(?![\d:])", r"\1:\2", t)
    return t


def seasons_of(s):
    s = s or ""
    if "모든" in s or not s.strip():
        return SEASONS[:]
    return [x for x in SEASONS if x in s]


def fish(sh):
    out = []

    def add(area, sub, name, season, weather, time, note, floor=None):
        name = name.strip()
        if not name:
            return
        out.append({
            "area": area, "sub": sub, "name": name,
            "seasons": seasons_of(season) if season is not None else SEASONS[:],
            "seasonText": season or "",
            "weather": weather or "", "time": fix_time(time or ""),
            "notes": lines(note), "floor": floor,
        })

    for r in range(6, 27):
        sub = "야시장(잠수함)" if r >= 24 else ""
        season = sh.val(r, "C")
        add("바다", sub, sh.raw(r, "B"), "겨울" if r >= 24 else season, sh.val(r, "D"), sh.val(r, "E"), sh.raw(r, "F"))
    for r in range(6, 22):
        add("강", sh.val(r, "H").replace(" ,", ","), sh.raw(r, "I"), sh.raw(r, "J"), sh.raw(r, "K"), sh.raw(r, "L"), sh.raw(r, "M"))
    for r in range(6, 12):
        add("산 호수", "", sh.raw(r, "O"), sh.raw(r, "P"), "", sh.raw(r, "Q"), sh.raw(r, "R"))
    for r in range(6, 16):
        add("게잡이 통발", sh.val(r, "T").strip(), sh.raw(r, "U"), None, "", "", sh.raw(r, "V"))
    for r in range(30, 36):
        add("그 외 장소", sh.val(r, "D").replace("\n", " "), sh.raw(r, "B"), None, "", sh.raw(r, "C"), sh.raw(r, "E"))
    for r in range(30, 40):
        add("진저섬", sh.val(r, "H"), sh.raw(r, "I"), None, sh.raw(r, "J"), sh.raw(r, "K"), sh.raw(r, "L"))
    for r in range(30, 36):
        add("광산", "", sh.raw(r, "O"), None, "", "", sh.raw(r, "Q"), floor=sh.raw(r, "P"))
    notes = [sh.raw(2, "B")] if sh.raw(2, "B") else []
    return {"list": out, "notes": notes, "video": sh.link(1, "B")}


# ---------------------------------------------------------------- 건물
def buildings(sh):
    robin, wizard = [], []
    for r in range(8, 28):
        name = sh.val(r, "B")
        if not name:
            continue
        tier = sh.raw(r, "C") if not sh.spans_cols(r, "B", "C") else ""
        robin.append({"name": name, "tier": tier, "cost": lines(sh.raw(r, "D")),
                      "note": lines(sh.raw(r, "E")), "days": sh.raw(r, "F")})
    for r in range(30, 40):
        name = sh.val(r, "B")
        if name:
            wizard.append({"name": name, "tier": "", "cost": lines(sh.raw(r, "D")),
                           "note": lines(sh.raw(r, "E")), "days": "즉시"})
    notes = [sh.raw(r, "A") for r in range(1, 6) if sh.raw(r, "A")]
    return {"robin": robin, "wizard": wizard, "notes": notes, "video": sh.link(4, "A")}


# ---------------------------------------------------------------- 요리
def ing_list(s):
    out = []
    for ln in lines(s):
        m = re.match(r"^(.*?)\s*\((\d+)\)$", ln)
        out.append({"name": m.group(1), "n": int(m.group(2))} if m else {"name": ln, "n": None})
    return out


def cooking(sh):
    recipes = []
    for r in range(5, sh.ws.max_row + 1):
        name = sh.raw(r, "B")
        if not name:
            continue
        recipes.append({"name": name, "ings": ing_list(sh.raw(r, "C")),
                        "how": lines(sh.raw(r, "D")), "note": lines(sh.raw(r, "E"))})
    # 오른쪽 "요리에 필요한 재료들만" 표. H:K 를 가로로 병합한 칸은 계절 무관.
    groups = []
    for r in range(5, 30):
        cat = sh.val(r, "G")
        if not cat:
            continue
        g = next((x for x in groups if x["cat"] == cat), None)
        if not g:
            g = {"cat": cat, "all": [], "seasons": {s: [] for s in SEASONS}}
            groups.append(g)
        if sh.spans_cols(r, "H", "K"):
            if sh.raw(r, "H"):
                g["all"] += lines(sh.raw(r, "H"))
        else:
            for s, c in zip(SEASONS, ("H", "I", "J", "K")):
                g["seasons"][s] += lines(sh.raw(r, c))
    notes = [sh.raw(r, "A") for r in range(1, 3) if sh.raw(r, "A")]
    return {"list": recipes, "ingredients": groups, "notes": notes}


# ---------------------------------------------------------------- 제작품
TITLE_FILL, HEAD_FILL = "FF134F5C", "FFD0E0E3"
# 상점 제작법 값이 위키(영어 「Carpenter's Shop」·「Crafting」, 한국어 아이템 문서)와 다른 것 — 위키 기준으로 고친다
CRAFT_FIX = {
    "나무바닥 (500골드)": ("나무바닥 (100골드)", None),
    "소박한판자바닥 (500골드)": ("소박한판자바닥 (200골드)", None),
    "돌바닥 (500골드)": ("돌바닥 (100골드)", None),
    "디딤돌길 (500골드)": ("디딤돌길 (100골드)", None),
    "짚바닥 (1,000골드)": ("짚바닥 (200골드)", None),
    "수정길 * 5 (1,000골드)": ("수정길 * 5 (200골드)", None),
    "큰석조보관함 (5,000골드)": ("큰석조보관함 (5,000골드)", "광산의 드워프에게서 구매"),  # 크로버스가 아니라 드워프
}


def crafting(sh):
    sections = []
    for cols in (("B", "C", "D"), ("F", "G", "H"), ("J", "K", "L")):
        cc, nc, mc = (ci(x) for x in cols)
        cur, cond = None, ""
        for r in range(33, sh.ws.max_row + 1):
            title = next((sh.raw(r, c) for c in (cc, nc) if sh.fill(r, c) == TITLE_FILL and sh.raw(r, c)), None)
            if title:
                cur = {"title": title, "condLabel": "", "items": [], "notes": [], "row": r, "col": cc}
                sections.append(cur)
                cond = ""
                continue
            if cur is None:
                continue
            if sh.fill(r, nc) == HEAD_FILL:
                cur["condLabel"] = sh.raw(r, cc) if sh.fill(r, cc) == HEAD_FILL else ""
                continue
            name, cval = sh.raw(r, nc), sh.raw(r, cc)
            if cval and not name and sh.spans_cols(r, cc, mc):
                cur["notes"].append(cval)
                continue
            if cval:
                cond = cval
            if name:
                fixed, where = CRAFT_FIX.get(name, (name, None))
                cur["items"].append({"cond": where or (cond if cur["condLabel"] or cond else ""), "name": fixed, "mats": sh.raw(r, mc)})
    sections.sort(key=lambda s: (s["row"], s["col"]))
    for s in sections:
        s.pop("row"), s.pop("col")
        s["title"] = s["title"].replace("배을", "배울")
        # 시트에 "전투" 가 두 번 적혀 있는데 오른쪽 칸(체리폭탄·계단…)은 채광 레시피다
        if s["title"].startswith("전투") and any(i["name"] == "체리폭탄" for i in s["items"]):
            s["title"] = s["title"].replace("전투", "채광")
    notes = [sh.raw(2, "A")] if sh.raw(2, "A") else []
    return {"sections": [s for s in sections if s["items"] or s["notes"]], "notes": notes, "video": sh.link(1, "A")}


# ---------------------------------------------------------------- 범위 그림
def grid(sh, r0, r1, c0, c1, legend):
    rows = []
    for r in range(r0, r1 + 1):
        rows.append("".join(legend.get(sh.fill(r, c), ".") for c in range(c0, c1 + 1)))
    return rows


def ranges(wb):
    bee = Sheet(wb["양봉장"])
    sc = Sheet(wb["허수아비"])
    sp = Sheet(wb["스프링클러 주니모오두막 범위"])
    L_BEE = {"FFF1C232": "B", "FF8E7CC3": "S", "FF6D9EEB": "C", "FFFF9DFF": "F"}
    L_SC = {"FFD9D2E9": "r", "FF6D9EEB": "C"}
    L_SP = {"FFCFE2F3": "r", "FF8E7CC3": "S", "FFEAD1DC": "r", "FFA4C2F4": "H"}
    return {
        "bee": {
            "notes": [bee.raw(r, "B") for r in range(2, 5) if bee.raw(r, "B")],
            "tips": lines(bee.raw(22, "B")),
            "grid": grid(bee, 8, 21, 2, 14, L_BEE),
            "video": bee.link(1, "B"),
        },
        "scarecrow": {
            "notes": [sc.raw(2, "A")],
            "items": [
                {"title": sc.raw(21, "F"), "grid": grid(sc, 4, 20, 2, 18, L_SC)},
                {"title": "허수아비 2개로 밭 덮기", "memo": sc.raw(22, "B"), "grid": grid(sc, 23, 31, 2, 18, L_SC)},
                {"title": sc.raw(66, "M"), "memo": sc.raw(67, "M"), "grid": grid(sc, 33, 65, 2, 34, L_SC)},
            ],
        },
        "sprinkler": {
            "notes": [sp.raw(2, "A")],
            "items": [
                {"title": sp.raw(4, "B"), "grid": grid(sp, 5, 7, 2, 4, L_SP)},
                {"title": sp.raw(4, "F"), "grid": grid(sp, 5, 7, 6, 8, L_SP)},
                {"title": sp.raw(4, "J"), "grid": grid(sp, 5, 9, 10, 14, L_SP)},
                {"title": sp.raw(4, "P"), "grid": grid(sp, 5, 11, 16, 22, L_SP)},
                {"title": sp.raw(32, "B"), "grid": grid(sp, 15, 31, 2, 18, L_SP)},
            ],
            "video": sp.link(1, "A"),
        },
    }


def main(src, dst):
    # data_only: 판매가·성장일 칸이 수식이라 계산된 값을 읽어야 한다
    wb = openpyxl.load_workbook(src, rich_text=True, data_only=True)
    vill = villagers(Sheet(wb["사랑하는 선물"]))
    data = {
        "source": "https://docs.google.com/spreadsheets/d/1y9wapWb8nqG0fVHTpCD82pj-Ak1Ms9VSbwV10oN5K5E/edit",
        "author": {"name": "에글이", "url": "https://www.youtube.com/@%EC%97%90%EA%B8%80%EC%9D%B4"},
        "basics": basics(Sheet(wb["기본 플레이"])),
        "calendar": calendar(Sheet(wb["달력"]), vill["list"]),
        "villagers": vill,
        "bundles": bundles(Sheet(wb["마을회관번들"])),
        "crops": crops(wb),
        "forage": forage(Sheet(wb["채집"])),
        "fish": fish(Sheet(wb["낚시"])),
        "ranges": ranges(wb),
        "buildings": buildings(Sheet(wb["건물 비용, 재료"])),
        "cooking": cooking(Sheet(wb["요리"])),
        "crafting": crafting(Sheet(wb["제작품"])),
    }
    js = "// tools/extract.py 가 만든 파일 — 손으로 고치지 말고 시트를 고친 뒤 다시 뽑는다.\nwindow.SDV = " + \
        json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    Path(dst).write_text(js, encoding="utf-8")
    print(f"wrote {dst} ({len(js) // 1024} KB)")
    print("  주민", len(vill["list"]), "· 작물", len(data["crops"]["list"]), "· 물고기", len(data["fish"]["list"]),
          "· 채집", len(data["forage"]["list"]), "· 요리", len(data["cooking"]["list"]),
          "· 제작 섹션", len(data["crafting"]["sections"]))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("사용법: python tools/extract.py 시트.xlsx [data/data.js]")
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else str(Path(__file__).resolve().parent.parent / "data" / "data.js"))
